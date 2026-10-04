"""
Put the project recap deck in front of the October deck, in one file.

    python scripts/merge_recap_deck.py

inputs  ../28:9:2026/Presentation/VIPP_201A_Project_Recap.pptx       (27 slides, 10 x 5.625 in)
        ../28:9:2026/Presentation/VIPP 301A Session Results October 2026.pptx   (18 slides, 13.33 x 7.5 in,
        from scripts/build_october_deck.py)
output  ../28:9:2026/Presentation/VIPP 301A Project Recap and October 2026 Update.pptx

order: recap 1 to 25, the 18 October slides, then the recap glossary (appendix) and the recap closing slide.
Recap slides are copied at the XML level with their images, charts, videos and timing; positions, sizes, line
widths, insets and font sizes are scaled by 4/3 (both decks are 16:9). Every footer becomes "N / 45".
"""
import copy
import os
import re

from lxml import etree
from pptx import Presentation
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI

HERE = os.path.dirname(os.path.abspath(__file__))
PRES = os.path.abspath(os.path.join(HERE, "..", "..", "28:9:2026", "Presentation"))
RECAP = os.path.join(PRES, "VIPP_201A_Project_Recap.pptx")
OCT = os.path.join(PRES, "VIPP 301A Session Results October 2026.pptx")
OUT = os.path.join(PRES, "VIPP 301A Project Recap and October 2026 Update.pptx")
OLD_URL = "github.com/xHeedz/VIPP-201---SPRING-2025-2026---Dr-Naseem-Daher"
NEW_URL = "github.com/xHeedz/VIPP-201-SPRING-2025-2026-Dr-Naseem-Daher"

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
R_ATTRS = ["{%s}%s" % (NS["r"], k) for k in ("id", "embed", "link", "pict")]
SKIP_RELS = ("slideLayout", "notesSlide")
GEOM_ATTRS = {"off": ("x", "y"), "ext": ("cx", "cy"), "chOff": ("x", "y"), "chExt": ("cx", "cy")}
EMU_ATTRS = ("w", "lIns", "rIns", "tIns", "bIns", "dist", "blurRad", "rad")


def scale_xml(root, k):
    """scale geometry, line widths, insets, shadows and font sizes of a slide or chart xml tree by k"""
    for el in root.iter():
        if not isinstance(el.tag, str):
            continue
        tag = etree.QName(el).localname
        if tag in GEOM_ATTRS and etree.QName(el).namespace in (NS["a"], NS["p"]):
            for at in GEOM_ATTRS[tag]:
                if at in el.attrib:
                    el.set(at, str(int(round(int(el.get(at)) * k))))
        for at in EMU_ATTRS:
            if at in el.attrib and el.get(at).lstrip("-").isdigit() and tag != "chOff":
                el.set(at, str(int(round(int(el.get(at)) * k))))
        if "sz" in el.attrib and tag in ("rPr", "defRPr", "endParaRPr") and el.get("sz").isdigit():
            el.set("sz", str(int(round(int(el.get("sz")) * k))))
        if tag == "spcPts" and "val" in el.attrib:
            el.set("val", str(int(round(int(el.get("val")) * k))))


def clone_part(src, pkg, cache):
    """copy a part (and, recursively, the parts it relates to) into pkg; returns the new part"""
    if src.partname in cache:
        return cache[src.partname]
    m = re.match(r"(.*?)(\d*)(\.\w+)$", str(src.partname))
    tmpl = m.group(1) + "%d" + m.group(3)
    new = Part(pkg.next_partname(tmpl), src.content_type, pkg, src.blob)
    cache[src.partname] = new
    mapping = {}
    for rId, rel in src.rels.items():
        if rel.is_external:
            mapping[rId] = new.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            mapping[rId] = new.rels.get_or_add(rel.reltype, clone_part(rel.target_part, pkg, cache))
    blob = src.blob
    if mapping and (src.content_type.endswith("+xml") or src.content_type.endswith("/xml")):
        root = etree.fromstring(blob)
        for el in root.iter():
            for at in R_ATTRS:
                if at in el.attrib and el.get(at) in mapping:
                    el.set(at, mapping[el.get(at)])
        if "drawingml.chart" in src.content_type:
            scale_xml(root, K)
        blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    elif "drawingml.chart" in src.content_type:
        root = etree.fromstring(blob)
        scale_xml(root, K)
        blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    new._blob = blob
    return new


def copy_slide(src_slide, dst_prs, cache):
    dst = dst_prs.slides.add_slide(dst_prs.slide_layouts[0])
    spart, dpart = src_slide.part, dst.part
    mapping = {}
    for rId, rel in spart.rels.items():
        if any(s in rel.reltype for s in SKIP_RELS):
            continue
        if rel.is_external:
            mapping[rId] = dpart.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            mapping[rId] = dpart.rels.get_or_add(rel.reltype, clone_part(rel.target_part, dst_prs.part.package, cache))
    src_el, dst_el = src_slide._element, dst._element
    for child in list(dst_el):
        dst_el.remove(child)
    for child in src_el:
        dst_el.append(copy.deepcopy(child))
    for el in dst_el.iter():
        for at in R_ATTRS:
            if at in el.attrib and el.get(at) in mapping:
                el.set(at, mapping[el.get(at)])
    scale_xml(dst_el, K)
    for t in dst_el.iter("{%s}t" % NS["a"]):
        if t.text and OLD_URL in t.text:
            t.text = t.text.replace(OLD_URL, NEW_URL)
    if src_slide.has_notes_slide and src_slide.notes_slide.notes_text_frame.text.strip():
        dst.notes_slide.notes_text_frame.text = src_slide.notes_slide.notes_text_frame.text
    return dst


def renumber(prs):
    """works on each slide part's xml: copied slides keep a cached shape tree in python-pptx that is detached"""
    total = len(prs.slides)
    for i, sl in enumerate(prs.slides, 1):
        for p in sl.part._element.iter("{%s}p" % NS["a"]):
            ts = list(p.iter("{%s}t" % NS["a"]))
            txt = "".join(t.text or "" for t in ts)
            if ts and re.fullmatch(r"\s*\d{1,2}\s*/\s*\d{1,2}\s*", txt):
                ts[0].text = f"{i} / {total}"
                for t in ts[1:]:
                    t.text = ""


def main():
    global K
    recap = Presentation(RECAP)
    prs = Presentation(OCT)
    K = prs.slide_width / recap.slide_width
    assert abs(K - prs.slide_height / recap.slide_height) < 1e-3, "aspect ratios differ"
    n_oct = len(prs.slides)
    cache = {}
    for s in recap.slides:
        copy_slide(s, prs, cache)
    ids = prs.slides._sldIdLst
    items = list(ids)
    oct_items, rec_items = items[:n_oct], items[n_oct:]
    order = rec_items[:25] + oct_items + rec_items[25:]
    for el in items:
        ids.remove(el)
    for el in order:
        ids.append(el)
    renumber(prs)
    prs.save(OUT)
    print("written", OUT, len(prs.slides), "slides, scale", round(K, 4))


if __name__ == "__main__":
    main()
