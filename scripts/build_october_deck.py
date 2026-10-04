"""
Append the October 2026 update to the September session deck, in the same style (Poppins / Source Serif 4,
AUB maroon, pink cards, footer with page numbers). Charts are native PowerPoint charts built from the result
files in data/, the planted driver animation is a GIF, the US-101 replay an embedded video.

    python scripts/build_october_deck.py
writes ../28:9:2026/Presentation/VIPP 301A Session Results October 2026.pptx (the September deck is not changed)
"""
import glob
import os
import re
import subprocess
import sys
import tempfile
import zipfile

import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..")
sys.path.insert(0, ROOT)
DATA = os.path.join(ROOT, "data")
FIG = os.path.join(ROOT, "results", "figures")
PRES = os.path.abspath(os.path.join(ROOT, "..", "28:9:2026", "Presentation"))
SRC = os.path.join(PRES, "VIPP 201A Session Results.pptx")
OUT = os.path.join(PRES, "VIPP 301A Session Results October 2026.pptx")
VIDEO = os.path.join(PRES, "us101_ellipse_replay.mp4")

# palette and fonts of the September deck
MAROON, MAROON2, WINE = "6E1423", "7A1F2B", "8C2A36"
INK, BODY, MUTED = "231A1C", "5E5456", "8E8385"
CARD, CARD_LINE, PINK, PALE, PAPER = "F5E9EB", "E9E2E3", "E8C9CE", "FBF3F4", "FBF8F6"
GREY_BAR, GREEN, YELLOW, RED = "B9AEB0", "2E9E4F", "E0B020", "C8323C"
HEAD, SANS = "Source Serif 4", "Poppins"
TOTAL = 18
FOOTER = "Hadi Al Shmaissani · VIPP 301A · Supervisor: Dr. Naseem Daher"
X0, W = 0.889, 11.556          # content left edge and width (inches)


def rgb(h):
    return RGBColor.from_string(h)


# ── drawing helpers ─────────────────────────────────────────────────────────
def text(slide, x, y, w, h, runs, size=12, color=BODY, font=SANS, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None, line=None, name=None):
    """runs: str, or list of paragraphs, each a str or a list of (text, {size,color,bold,font}) runs"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line:
            p.line_spacing = line
        if i > 0:
            p.space_before = Pt(4)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = rgb(o.get("color", color))
            if spacing:
                rPr = r._r.get_or_add_rPr()
                rPr.set("spc", str(spacing))
    return tb


def card(slide, x, y, w, h, fill=CARD, line=None, radius=0.06, name=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.5)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if name:
        s.name = name
    return s


def dot(slide, x, y, d, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def new_slide(prs, bg=PAPER):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(bg)
    return s


def header(slide, label, title, crest):
    text(slide, X0, 0.722, W, 0.26, label.upper(), size=12, color=WINE, spacing=150)
    text(slide, X0, 0.997, 10.4, 0.55, [[(title, {"font": HEAD, "size": 30, "color": INK})]], line=1.0)
    slide.shapes.add_picture(crest, Inches(12.014), Inches(0.611), Inches(0.431), Inches(0.479))


def footer(slide, n):
    text(slide, X0, 6.822, 9.3, 0.26, FOOTER, size=12, color=MUTED)
    text(slide, 11.556, 6.822, 0.889, 0.26, f"{n} / {TOTAL}", size=12, color=MUTED, align=PP_ALIGN.RIGHT)


def style_chart(ch, legend=True, size=10):
    ch.font.name = SANS
    ch.font.size = Pt(size)
    ch.font.color.rgb = rgb(BODY)
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size)
    try:
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = rgb("ECE6E7")
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(size)
        ca = ch.category_axis
        ca.format.line.color.rgb = rgb("D9D0D2")
        ca.tick_labels.font.size = Pt(size)
        ca.has_major_gridlines = False
    except Exception:
        pass


def color_series(ch, colors):
    for s, c in zip(ch.plots[0].series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(c)
        s.format.line.fill.background()


def chart_card(slide, x, y, w, h, title=None):
    card(slide, x, y, w, h, fill="FFFFFF", line=CARD_LINE, radius=0.04)
    if title:
        text(slide, x + 0.2, y + 0.14, w - 0.4, 0.3, [[(title, {"bold": True, "color": INK})]], size=12)


def stat(slide, x, y, w, big, small, sub, dark=False):
    card(slide, x, y, w, 0.95, fill=MAROON if dark else "FFFFFF", line=None if dark else CARD_LINE)
    text(slide, x + 0.18, y + 0.13, 2.3, 0.6, [[(big, {"size": 26, "bold": True, "color": PALE if dark else MAROON2}),
                                                (small, {"size": 26, "bold": True, "color": PINK if dark else MUTED})]],
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + 2.55, y + 0.12, w - 2.7, 0.72, sub, size=12, color=PALE if dark else INK, anchor=MSO_ANCHOR.MIDDLE)


def bullet_lines(slide, x, y, w, items, color=INK, marker=MAROON2, size=12, step=0.5):
    for i, it in enumerate(items):
        dot(slide, x, y + i * step + 0.07, 0.1, marker)
        text(slide, x + 0.22, y + i * step, w - 0.22, step, it, size=size, color=color)


# ── data ────────────────────────────────────────────────────────────────────
def hist_share(scores, edges):
    h, _ = np.histogram(np.clip(scores, 0, 99.999), bins=edges)
    return 100 * h / max(h.sum(), 1)


def calibration_curves():
    from datasets.ngsim import read_raw, trajectories
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        tr = trajectories(read_raw(os.path.join(ROOT, "..", "28:9:2026", "ngsim_us-101_5min.csv"), "us-101"))
    tr = tr[np.isclose(tr["t"] % 1.0, 0.0, atol=0.05)]
    m = (tr["gap"] > 0) & (tr["speed"] > 1)
    ref = (tr["gap"] / tr["speed"])[m].to_numpy()
    sims = []
    for f in sorted(glob.glob(os.path.join(DATA, "sumo_planted", "jam_jam_s*.per_second.csv.gz"))):
        ps = pd.read_csv(f, usecols=["speed_kmh", "gap_m"])
        mm = (ps["gap_m"] > 0) & (ps["speed_kmh"] > 3.6)
        sims.append((ps["gap_m"] / (ps["speed_kmh"] / 3.6))[mm].to_numpy())
    sim = np.concatenate(sims) if sims else np.array([])
    edges = np.arange(0, 4.01, 0.25)

    def share(x):          # rows above 4 s are left out (not piled into the last bin)
        h, _ = np.histogram(x[x < edges[-1]], bins=edges)
        return 100 * h / max(len(x), 1)
    return edges, share(ref), share(sim)


def poster(video):
    d = tempfile.mkdtemp()
    subprocess.run(["qlmanage", "-t", "-s", "1600", "-o", d, video], capture_output=True)
    pngs = glob.glob(os.path.join(d, "*.png"))
    return pngs[0] if pngs else None


# ── slides ──────────────────────────────────────────────────────────────────
def divider(prs, logo):
    s = new_slide(prs, MAROON)
    s.shapes.add_picture(logo, Inches(X0), Inches(0.667), Inches(3.333), Inches(0.778))
    text(s, 8.2, 0.6, 4.244, 1.0, [[("Hadi Al Shmaissani", {"bold": True, "color": "FFFFFF"})],
                                    [("Supervisor: Dr. Naseem Daher", {"color": PALE})],
                                    [("October 2026 · American University of Beirut", {"color": PALE})]],
         size=12, align=PP_ALIGN.RIGHT)
    text(s, X0, 1.85, W, 0.3, "VIPP 301A · OCTOBER 2026 UPDATE", size=12, color=PINK, spacing=150)
    text(s, X0, 2.2, W, 0.8, [[("Hard-checking every number", {"font": HEAD, "size": 40, "color": "FFFFFF"})]])
    text(s, X0, 3.05, 8.5, 0.7, "Dr. Daher's feedback after the September session: verify and hand check everything, "
         "run better simulations, train on more data, then give the model context", size=16, color=PALE)
    steps = [("1 · Verify", "three SUMO bugs fixed, every input hand checked"),
             ("2 · Simulate", "planted drivers with true labels, calibrated to US-101"),
             ("3 · Datasets", "NGSIM arterials, I-80 and pNEUMA Athens"),
             ("4 · Context", "retrieve normal drivers in the same situation")]
    w = (W - 3 * 0.3) / 4
    for i, (a, b) in enumerate(steps):
        x = X0 + i * (w + 0.3)
        card(s, x, 4.45, w, 1.35, fill=WINE)
        text(s, x + 0.2, 4.62, w - 0.4, 0.3, [[(a, {"bold": True, "color": "FFFFFF", "size": 14})]])
        text(s, x + 0.2, 5.0, w - 0.4, 0.7, b, size=12, color=PALE)
    s.notes_slide.notes_text_frame.text = ("Four parts, done in order. Measurement first, context last: "
                                           "a context layer on an unverified index would only hide bugs.")
    return s


def slide_bugs(prs, crest, n):
    s = new_slide(prs)
    header(s, "Part 1 · Verify", "Three measurement bugs inflated the SUMO scores", crest)
    bugs = [("Lane offset read travel distance",
             "std of the global y position over 30 samples: a car driving north got 40 points from travel alone. "
             "Now the offset from the lane centre."),
            ("Acceleration 3 to 4 times too large",
             "speed change divided by 0.1 s, sampled every 0.3 to 0.4 s. Now SUMO's own acceleration."),
            ("Proximity measured to the ego car",
             "straight-line distance to the RL car, not the gap to the car ahead. Now the gap to the leader, "
             "as in UAH and NGSIM.")]
    for i, (a, b) in enumerate(bugs):
        y = 1.65 + i * 1.32
        card(s, X0, y, 5.25, 1.17)
        text(s, X0 + 0.2, y + 0.14, 4.85, 0.3, [[(a, {"bold": True, "color": MAROON2, "size": 13})]])
        text(s, X0 + 0.2, y + 0.46, 4.85, 0.68, b, size=12)
    before = pd.read_csv(os.path.join(DATA, "sumo_npc_aggressiveness_before_fix.csv"))["ai_score"]
    after = pd.read_csv(os.path.join(DATA, "sumo_npc_aggressiveness.csv"))["ai_score"]
    edges = np.arange(0, 101, 10)
    cd = CategoryChartData()
    cd.categories = [f"{a}-{a + 10}" for a in edges[:-1]]
    cd.add_series(f"before the fixes (median {before.median():.1f})", [round(v, 1) for v in hist_share(before, edges)])
    cd.add_series(f"after (median {after.median():.1f})", [round(v, 1) for v in hist_share(after, edges)])
    chart_card(s, 6.45, 1.65, 5.995, 3.81, "SUMO urban NPCs: share of rows per score band (%)")
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.6), Inches(2.05), Inches(5.7), Inches(3.3), cd)
    ch = gf.chart
    style_chart(ch)
    color_series(ch, [GREY_BAR, MAROON2])
    ch.plots[0].gap_width = 60
    ch.value_axis.maximum_scale = 100
    stat(s, 6.45, 5.62 - 0.05, 5.995, f"{before.median():.1f}", f"  to {after.median():.1f}",
         "median urban score: 99.4% of rows had saturated waviness before the fix", dark=True)
    footer(s, n)
    s.notes_slide.notes_text_frame.text = ("The same three bugs were in three files: both SUMO agent scripts and "
                                           "collect_data.py. Expected urban median drop about 40 points; measured 40.")
    return s


def slide_hand(prs, crest, n):
    s = new_slide(prs)
    header(s, "Part 1 · Verify", "Every input re-derived by hand matches the code", crest)
    checks = [("UAH normal", "42.32", "motorway, 107 km/h, 30 m gap"),
              ("UAH aggressive", "86.87", "motorway, 125 km/h, 12 m gap"),
              ("NGSIM US-101", "52.50", "car 983, 34 km/h, 15.8 m gap"),
              ("NGSIM Lankershim", "77.12", "arterial, 13 km/h, 4.6 m gap"),
              ("pNEUMA Athens", "57.19", "car 19 behind a stopped car")]
    w = (W - 4 * 0.2) / 5
    for i, (a, b, c) in enumerate(checks):
        x = X0 + i * (w + 0.2)
        card(s, x, 1.65, w, 1.75, fill="FFFFFF", line=CARD_LINE)
        text(s, x + 0.18, 1.8, w - 0.36, 0.3, [[(a, {"bold": True, "color": MAROON2})]], size=12)
        text(s, x + 0.18, 2.12, w - 0.36, 0.55, [[(b, {"size": 28, "bold": True, "color": INK})]])
        text(s, x + 0.18, 2.72, w - 0.36, 0.3, "hand = code", size=12, color=GREEN, bold=True)
        text(s, x + 0.18, 2.98, w - 0.36, 0.4, c, size=10, color=MUTED)
    card(s, X0, 3.62, 5.6, 2.55)
    text(s, X0 + 0.25, 3.8, 5.1, 0.3, [[("One formula instead of four", {"bold": True, "color": MAROON2, "size": 14})]])
    text(s, X0 + 0.25, 4.2, 5.1, 2.3, ["Four different definitions were in use: the SUMO validation runner, the "
                                       "highway-env assessors, the SUMO NPC scripts and the UAH agent.",
                                       "Every script now imports one reference, model/aggressiveness_model.py, and "
                                       "SUMO inputs come from one function. Hand numbers are hard coded in the tests."],
         size=12, color=INK)
    x2 = X0 + 5.85
    card(s, x2, 3.62, W - 5.85, 2.55)
    text(s, x2 + 0.25, 3.8, 5.2, 0.3, [[("What the hand checks caught", {"bold": True, "color": MAROON2, "size": 14})]])
    bullet_lines(s, x2 + 0.25, 4.25, W - 6.35, [
        "NGSIM reuses vehicle IDs in each recording period: two cars were spliced into one trajectory",
        "Arterials curve: road curvature was read as weaving (24% to 2% of rows above 1.5 m)",
        "UAH used lane rows where the lane was not detected (12% of rows)",
        "Legacy dataset: acceleration 15 times too large (wrong time step)"], step=0.56)
    footer(s, n)
    s.notes_slide.notes_text_frame.text = ("Each check recomputes every input from the raw files with plain Python, "
                                           "no pipeline code, and agrees with the pipeline to about 1e-11.")
    return s


def slide_high(prs, crest, n):
    s = new_slide(prs)
    header(s, "Part 1 · Verify", "High scores come from proximity measured in metres", crest)
    t = pd.read_csv(os.path.join(DATA, "top5_contributions.csv"))
    pick = [("uah", "all rows", "UAH, all windows"), ("uah", "top 5%", "UAH, top 5%"),
            ("ngsim us-101", "all rows", "US-101, all rows"), ("ngsim us-101", "top 5%", "US-101, top 5%")]
    cd = CategoryChartData()
    cd.categories = [p[2] for p in pick]
    rows = [t[(t["source"] == a) & (t["rows"] == b)].iloc[0] for a, b, _ in pick]
    for k, name in [("speed", "speed"), ("accel", "acceleration"), ("prox", "proximity"), ("wave", "lane offset")]:
        cd.add_series(name, [round(float(r[k]), 1) for r in rows])
    chart_card(s, X0, 1.65, 6.6, 4.92, "Points per term (mean)")
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(X0 + 0.15), Inches(2.05), Inches(6.3), Inches(4.4), cd)
    ch = gf.chart
    style_chart(ch)
    color_series(ch, [GREY_BAR, PINK, MAROON2, "D9A441"])
    ch.plots[0].gap_width = 55
    ch.category_axis.reverse_order = True
    pl = ch.plots[0]
    pl.has_data_labels = True
    pl.data_labels.number_format = "0"
    pl.data_labels.number_format_is_linked = False
    pl.data_labels.position = XL_LABEL_POSITION.CENTER
    pl.data_labels.font.size = Pt(10)
    pl.data_labels.font.color.rgb = rgb("FFFFFF")
    ng = rows[3]
    x2 = X0 + 6.85
    stat(s, x2, 1.65, W - 6.85, f"{ng['prox']:.0f}", f" of {ng['score']:.0f}",
         f"points from proximity in the US-101 top 5%; speed gives {ng['speed']:.0f}")
    stat(s, x2, 2.8, W - 6.85, "1.7 s", "", "headway of US-101 car 983 (15.8 m at 34 km/h): ordinary, "
         "yet 37 proximity points")
    card(s, x2, 3.95, W - 6.85, 2.62, fill=MAROON)
    text(s, x2 + 0.25, 4.15, W - 7.35, 0.3, [[("Metres treat a slow queue as tailgating", {"bold": True, "color": "FFFFFF", "size": 14})]])
    text(s, x2 + 0.25, 4.6, W - 7.35, 1.9, "A 10 m gap gives the same proximity term at 20 km/h and at 120 km/h. "
         "Time headway (gap divided by speed) separates them: 1.8 s in a queue, 0.3 s at motorway speed.",
         size=12, color=PALE)
    footer(s, n)
    return s


def slide_planted(prs, crest, n):
    s = new_slide(prs)
    header(s, "Part 2 · Simulate", "Simulated drivers with known labels", crest)
    gif = os.path.join(FIG, "planted_jam_animation.gif")
    card(s, X0, 1.65, 8.1, 3.75, fill="2B2B2B", radius=0.03)
    s.shapes.add_picture(gif, Inches(X0 + 0.08), Inches(1.73), Inches(7.94), Inches(3.475))
    text(s, X0, 5.52, 8.1, 0.95, "Calibrated jam before a lane drop, the same cars three times. Top: the true driver type. "
         "Middle: the index with the gap in metres. Bottom: with time headway. Green conservative, yellow normal, "
         "red aggressive.", size=12)
    x2 = X0 + 8.35
    w2 = W - 8.35
    text(s, x2, 1.65, w2, 0.3, [[("Three planted driver types (IDM)", {"bold": True, "color": INK})]], size=12)
    types = [("conservative", GREEN, "1.26 s headway, gentle acceleration"),
             ("normal", YELLOW, "0.84 s, the SUMO default range"),
             ("aggressive", RED, "0.42 s, hard accel and braking, fast")]
    for i, (a, c, b) in enumerate(types):
        y = 2.05 + i * 0.68
        dot(s, x2, y + 0.06, 0.16, c)
        text(s, x2 + 0.28, y, w2 - 0.28, 0.6, [[(a, {"bold": True, "color": INK})], [(b, {"size": 10, "color": BODY})]], size=12)
    text(s, x2, 4.15, w2, 0.3, [[("True normal drivers in the jam", {"bold": True, "color": INK})]], size=12)
    text(s, x2, 4.45, w2, 0.3, "labelled aggressive:", size=12)
    card(s, x2, 4.85, w2 / 2 - 0.08, 1.05, fill="FFFFFF", line=CARD_LINE)
    text(s, x2 + 0.12, 4.95, w2 / 2 - 0.3, 0.9, [[("65%", {"size": 26, "bold": True, "color": RED})], [("metres", {"color": MUTED})]], size=12)
    card(s, x2 + w2 / 2 + 0.08, 4.85, w2 / 2 - 0.08, 1.05, fill="FFFFFF", line=CARD_LINE)
    text(s, x2 + w2 / 2 + 0.2, 4.95, w2 / 2 - 0.3, 0.9, [[("31%", {"size": 26, "bold": True, "color": MAROON2})], [("headway", {"color": MUTED})]], size=12)
    text(s, x2, 6.0, w2, 0.5, "mix 20 / 60 / 20; 7 scenarios x 10 seeds; label in the vehicle id", size=10, color=MUTED)
    footer(s, n)
    s.notes_slide.notes_text_frame.text = ("Driver types differ by IDM tau, acceleration, deceleration, desired speed "
                                           "and lane change assertiveness; lateral wobble is the same for all three.")
    return s


def slide_planted_results(prs, crest, n):
    s = new_slide(prs)
    header(s, "Part 2 · Simulate", "Good ranking; headway stops the density climb", crest)
    sm = pd.read_csv(os.path.join(DATA, "planted_summary.csv"))
    order = [("highway", "low", "free"), ("highway", "medium", "medium"), ("jam", "jam", "jam"), ("merge", "medium", "merge"),
             ("urban", "medium", "junction"), ("roundabout", "medium", "roundabout"), ("weather", "medium", "rain")]
    rows = [sm[(sm["scenario"] == a) & (sm["density"] == b)].iloc[0] for a, b, _ in order]
    cd = CategoryChartData()
    cd.categories = [o[2] for o in order]
    cd.add_series("index, metres", [round(float(r["auc_reference"]), 3) for r in rows])
    cd.add_series("index, headway", [round(float(r["auc_headway"]), 3) for r in rows])
    cd.add_series("UAH-trained agent", [round(float(r["auc_agent"]), 3) for r in rows])
    chart_card(s, X0, 1.65, 6.3, 4.92, "AUC, aggressive vs normal (mean of 10 seeds)")
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(X0 + 0.15), Inches(2.05), Inches(6.0), Inches(4.4), cd).chart
    style_chart(ch)
    color_series(ch, [MAROON2, "D9A441", GREY_BAR])
    ch.value_axis.minimum_scale = 0.5
    ch.value_axis.maximum_scale = 1.0
    ch.value_axis.tick_labels.number_format = "0.00"
    ch.value_axis.tick_labels.number_format_is_linked = False
    ch.value_axis.major_unit = 0.1
    ch.plots[0].gap_width = 70
    d = pd.read_csv(os.path.join(DATA, "planted_density_headway.csv"), header=[0, 1], index_col=[0, 1])
    cd2 = CategoryChartData()
    cd2.categories = ["free", "medium", "jam"]
    for prox, typ, nm in [("metres", "normal", "normal drivers, metres"), ("headway", "normal", "normal drivers, headway"),
                          ("metres", "aggressive", "aggressive, metres"), ("headway", "aggressive", "aggressive, headway")]:
        r = d.loc[(prox, typ)]
        cd2.add_series(nm, [round(float(r[("median", k)]), 1) for k in ("low", "medium", "jam")])
    x2 = X0 + 6.55
    chart_card(s, x2, 1.65, W - 6.55, 4.92, "Median score, same highway, rising density")
    ch2 = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x2 + 0.15), Inches(2.05), Inches(W - 6.85), Inches(4.4), cd2).chart
    style_chart(ch2)
    ch2.value_axis.minimum_scale = 0
    ch2.value_axis.maximum_scale = 80
    ch2.value_axis.major_unit = 20
    for ser, c, dash in zip(ch2.plots[0].series, [MAROON2, MAROON2, "D9A441", "D9A441"], [False, True, False, True]):
        ser.format.line.color.rgb = rgb(c)
        ser.format.line.width = Pt(2.25)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            ser.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        ser.marker.format.fill.solid()
        ser.marker.format.fill.fore_color.rgb = rgb(c)
        ser.smooth = False
    footer(s, n)
    s.notes_slide.notes_text_frame.text = ("In metres the same normal drivers gain about 17 points from free flow to jam, "
                                           "almost all from proximity; with headway they stay flat.")
    return s


def slide_calibration(prs, crest, n):
    s = new_slide(prs)
    header(s, "Part 2 · Simulate", "Calibrated headways and a safer ego car", crest)
    edges, ref, sim = calibration_curves()
    cd = CategoryChartData()
    cd.categories = [f"{a:.2f}" for a in edges[:-1]]
    cd.add_series("NGSIM US-101", [round(v, 2) for v in ref])
    cd.add_series("SUMO jam, calibrated drivers", [round(v, 2) for v in sim])
    chart_card(s, X0, 1.65, 6.3, 3.55, "Time headway (s): share of rows per 0.25 s (%)")
    ch = s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(X0 + 0.15), Inches(2.05), Inches(6.0), Inches(3.05), cd).chart
    style_chart(ch)
    for ser, c in zip(ch.plots[0].series, [INK, MAROON2]):
        ser.format.line.color.rgb = rgb(c)
        ser.format.line.width = Pt(2.5)
        ser.smooth = True
    ch.category_axis.tick_labels.font.size = Pt(9)
    ch.value_axis.minimum_scale = 0
    ch.value_axis.maximum_scale = 25
    ch.value_axis.major_unit = 5
    cal = pd.read_csv(os.path.join(DATA, "planted_calibration.csv"))
    old = pd.read_csv(os.path.join(DATA, "planted_calibration_tau1.csv"))
    ks_new = float(cal[(cal["density"] == "jam") & (cal["signal"] == "time_headway_s")]["ks_d"].iloc[0])
    ks_old = float(old[(old["density"] == "jam") & (old["signal"] == "time_headway_s")]["ks_d"].iloc[0])
    stat(s, X0, 5.35, 6.3, f"{ks_old:.2f}", f"  to {ks_new:.2f}", "KS distance of the jam headways after tau x 0.7 "
         "(1.26 / 0.84 / 0.42 s). Speeds still differ: see notes")
    x2 = X0 + 6.55
    w2 = W - 6.55
    text(s, x2, 1.65, w2, 0.3, [[("RL ego car", {"bold": True, "color": INK, "size": 14})]])
    ego = [("30 m/s2", "  to 0", "ego emergency braking: a brake step was applied in 0.1 s; now limited to 5 m/s2"),
           ("176", "  to 1", "highway collisions: lane changes now respect the safe gaps of other cars"),
           ("4 of 200", "", "urban episodes with a junction crash, mostly early exploration; some are NPCs "
                            "driving into the ego")]
    for i, (a, b, c) in enumerate(ego):
        stat(s, x2, 2.05 + i * 1.1, w2, a, b, c, dark=(i == 0))
    text(s, x2, 5.4, w2, 1.0, "Full SUMO safety removes the last crashes but the ego then stalls in 43% of episodes: "
         "the fix belongs in the RL reward.", size=12, color=BODY)
    us = os.path.join(DATA, "us101_sweep.csv")
    note = "US-101 layout (5 lanes, ramps, auxiliary lane): see us101_sweep.csv."
    if os.path.exists(us):
        u = pd.read_csv(us).sort_values("ks_sum").iloc[0]
        note = (f"US-101 like layout, best demand {int(u['main'])} + {int(u['ramp'])} veh/h: median speed "
                f"{u['speed_med']:.0f} km/h (US-101 48), headway KS {u['ks_thw']:.2f}, speed KS {u['ks_speed']:.2f}.")
    s.notes_slide.notes_text_frame.text = ("Calibration sweep: tau multiplier 0.5 to 1.0, two demands, two seeds. " + note)
    footer(s, n)
    return s


def slide_real(prs, crest, n):
    s = new_slide(prs)
    header(s, "Part 3 · Datasets", "On real traffic, metres reads congestion as aggression", crest)
    u = pd.read_csv(os.path.join(DATA, "unlabelled_summary.csv")).set_index("site")
    sites = [("uah normal", "UAH normal"), ("uah aggressive", "UAH aggressive"), ("ngsim us-101", "US-101"),
             ("ngsim i-80", "I-80"), ("ngsim lankershim", "Lankershim"), ("ngsim peachtree", "Peachtree"),
             ("pneuma athens", "Athens (pNEUMA)")]
    sites = [p for p in sites if p[0] in u.index]
    cd = CategoryChartData()
    cd.categories = [p[1] for p in sites]
    cd.add_series("gap in metres", [round(float(u.loc[p[0], "score_median"]), 1) for p in sites])
    cd.add_series("time headway", [round(float(u.loc[p[0], "headway_score_median"]), 1) for p in sites])
    chart_card(s, X0, 1.65, 6.3, 4.92, "Median window score")
    ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(X0 + 0.15), Inches(2.05), Inches(6.0), Inches(4.4), cd).chart
    style_chart(ch)
    color_series(ch, [MAROON2, "D9A441"])
    ch.category_axis.reverse_order = True
    ch.value_axis.minimum_scale = 0
    ch.value_axis.maximum_scale = 80
    ch.value_axis.major_unit = 20
    ch.plots[0].gap_width = 45
    pl = ch.plots[0]
    pl.has_data_labels = True
    pl.data_labels.number_format = "0"
    pl.data_labels.number_format_is_linked = False
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    pl.data_labels.font.size = Pt(9)
    d = pd.read_csv(os.path.join(DATA, "unlabelled_density.csv"))
    bins = ["<10", "10-20", "20-40", "40-60", ">60"]
    cd2 = CategoryChartData()
    cd2.categories = bins
    for road, prox, col, nm in [("freeway", "score", MAROON2, "freeway, metres"), ("freeway", "score_headway", MAROON2, "freeway, headway"),
                                ("arterial", "score", GREEN, "arterial, metres"), ("arterial", "score_headway", GREEN, "arterial, headway")]:
        g = d[d["road_type"] == road].set_index("density_bin")
        cd2.add_series(nm, [round(float(g.loc[b, prox]), 1) if b in g.index else None for b in bins])
    x2 = X0 + 6.55
    chart_card(s, x2, 1.65, W - 6.55, 4.92, "NGSIM: median score by density (veh/km/lane)")
    ch2 = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x2 + 0.15), Inches(2.05), Inches(W - 6.85), Inches(4.4), cd2).chart
    style_chart(ch2)
    ch2.value_axis.minimum_scale = 0
    ch2.value_axis.maximum_scale = 80
    ch2.value_axis.major_unit = 20
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    for ser, c, dash in zip(ch2.plots[0].series, [MAROON2, MAROON2, GREEN, GREEN], [False, True, False, True]):
        ser.format.line.color.rgb = rgb(c)
        ser.format.line.width = Pt(2.25)
        if dash:
            ser.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        ser.marker.format.fill.solid()
        ser.marker.format.fill.fore_color.rgb = rgb(c)
    footer(s, n)
    s.notes_slide.notes_text_frame.text = ("Congested I-80 traffic at 26 km/h scores above the UAH aggressive drivers "
                                           "in metres. With time headway every unlabelled site sits near UAH normal.")
    return s


def slide_labels(prs, crest, n):
    s = new_slide(prs)
    header(s, "Part 3 · Datasets", "Looking for labelled dense traffic", crest)
    rows = [("100-DrivingStyle (2024)", "expert + self rated aggressiveness, 100 drivers", "speed, accel, pedals; no gap", "email to authors", YELLOW),
            ("DriveDNA (2026)", "rule-based primitives, 465 drivers", "all index signals, stop and go", "sample in use", GREEN),
            ("POLIDriving (Quito)", "accident risk level", "speed, accel; heavy traffic", "open", GREY_BAR),
            ("CAN-DSAD", "aggressive braking / lane change events", "CAN bus", "account needed", GREY_BAR),
            ("highD / exiD, SHRP2", "none / crash events", "full trajectories", "skipped", GREY_BAR)]
    cols = [(X0 + 0.2, 2.9, "dataset"), (X0 + 3.2, 3.4, "labels"), (X0 + 6.7, 2.7, "signals"), (X0 + 9.5, 1.9, "status")]
    for x, w, h in cols:
        text(s, x, 1.65, w, 0.3, h.upper(), size=10, color=WINE, spacing=120)
    for i, (a, b, c, d, col) in enumerate(rows):
        y = 2.0 + i * 0.62
        card(s, X0, y, W, 0.52, fill="FFFFFF" if i % 2 else CARD, line=None)
        text(s, cols[0][0], y + 0.14, cols[0][1], 0.3, [[(a, {"bold": True, "color": INK})]], size=12)
        text(s, cols[1][0], y + 0.14, cols[1][1], 0.3, b, size=12)
        text(s, cols[2][0], y + 0.14, cols[2][1], 0.3, c, size=12)
        dot(s, cols[3][0], y + 0.2, 0.13, col)
        text(s, cols[3][0] + 0.22, y + 0.14, cols[3][1], 0.3, d, size=12, color=INK)
    card(s, X0, 5.2, 7.0, 1.37, fill=MAROON)
    text(s, X0 + 0.25, 5.36, 6.5, 0.3, [[("Chosen: 100-DrivingStyle + DriveDNA", {"bold": True, "color": "FFFFFF", "size": 14})]])
    text(s, X0 + 0.25, 5.74, 6.5, 0.8, "Human ratings for the speed and acceleration terms; every term, including "
         "headway, in stop-and-go traffic with 465 drivers.", size=12, color=PALE)
    card(s, X0 + 7.25, 5.2, W - 7.25, 1.37, fill="FFFFFF", line=CARD_LINE)
    r = pd.read_csv(os.path.join(DATA, "uah_relabel.csv"))
    g = r[r["proximity"] == "metres"].groupby("labels")["agent_auc"].mean()
    text(s, X0 + 7.45, 5.36, W - 7.65, 0.3, [[("UAH per-minute labels", {"bold": True, "color": MAROON2})]], size=12)
    text(s, X0 + 7.45, 5.7, W - 7.65, 0.85, f"only 47% of aggressive-trip windows look aggressive; training on them: "
         f"agent AUC {g['trip']:.2f} to {g['window']:.2f}", size=12, color=INK)
    footer(s, n)
    return s


def slide_drivedna(prs, crest, n):
    s = new_slide(prs)
    header(s, "Part 3 · Datasets", "Radar car following: metres misreads queues", crest)
    sm = pd.read_csv(os.path.join(DATA, "drivedna_summary.csv")).set_index("question")
    fol = [q for q in sm.index if q.startswith("share labelled aggressive") and "slow car following" in q][0]
    fre = [q for q in sm.index if q.startswith("share labelled aggressive") and "free driving" in q][0]
    cd = CategoryChartData()
    cd.categories = ["slow car following (< 30 km/h, radar leader)", "free driving (>= 60 km/h, no leader)"]
    cd.add_series("gap in metres", [round(100 * float(sm.loc[fol, "metres"]), 1), round(100 * float(sm.loc[fre, "metres"]), 1)])
    cd.add_series("time headway", [round(100 * float(sm.loc[fol, "headway"]), 1), round(100 * float(sm.loc[fre, "headway"]), 1)])
    chart_card(s, X0, 1.65, 6.3, 4.92, "Windows labelled aggressive (%), DriveDNA sample")
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(X0 + 0.15), Inches(2.05), Inches(6.0), Inches(4.4), cd).chart
    style_chart(ch)
    color_series(ch, [MAROON2, "D9A441"])
    ch.value_axis.minimum_scale = 0
    ch.value_axis.maximum_scale = 100
    ch.value_axis.major_unit = 20
    pl = ch.plots[0]
    pl.gap_width = 70
    pl.has_data_labels = True
    pl.data_labels.number_format = "0.0"
    pl.data_labels.number_format_is_linked = False
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    pl.data_labels.font.size = Pt(10)
    x2 = X0 + 6.55
    w2 = W - 6.55
    w = pd.read_csv(os.path.join(DATA, "drivedna_windows.csv.gz"))
    text(s, x2, 1.65, w2, 1.2, [[("DriveDNA sample", {"bold": True, "color": INK, "size": 14})],
                                 f"{w['trip'].nunique()} real drives, {w['driver'].nunique()} drivers, {w['car'].nunique()} cars, "
                                 "openpilot logs at 10 Hz: radar gap to the lead car, lane offsets, speed. Human driving only."],
         size=12)
    stat(s, x2, 2.95, w2, f"{float(sm.loc[[q for q in sm.index if q.startswith('median score, slow')][0], 'metres']):.0f}",
         f" vs {float(sm.loc[[q for q in sm.index if q.startswith('median score, slow')][0], 'headway']):.0f}",
         "median score in slow car following, metres vs headway")
    sd = sm.loc[[q for q in sm.index if q.startswith("driver_078, spread")][0]]
    stat(s, x2, 4.05, w2, f"{float(sd['metres']):.0f} pts", "", "spread of one driver's median score across 13 cars: "
         "the car and route still move the score")
    text(s, x2, 5.2, w2, 1.3, "No aggressiveness labels: this tests the measurement, not the label. Lane positions are "
         "missing for 8 car models, so the score here uses speed, acceleration and proximity.", size=12, color=BODY)
    footer(s, n)
    s.notes_slide.notes_text_frame.text = ("DriveDNA-Sample (HenryYHW/DriveDNA-Sample, research-only licence). The full "
                                           "DriveDNA release (465 drivers) is waiting for the authors' approval.")
    return s


def slide_rag(prs, crest, n):
    s = new_slide(prs)
    header(s, "Part 4 · Context", "Context helps only where situations are mixed", crest)
    card(s, X0, 1.65, 6.5, 3.0, fill="FFFFFF", line=CARD_LINE, radius=0.04)
    s.shapes.add_picture(os.path.join(FIG, "rag_design.png"), Inches(X0 + 0.12), Inches(1.75), Inches(6.26), Inches(2.39))
    text(s, X0 + 0.2, 4.2, 6.1, 0.4, "Neighbours are chosen by situation, not by driving: the score is the percentile "
         "among normal drivers in the same situation.", size=10, color=MUTED)
    r = pd.read_csv(os.path.join(DATA, "rag_ablation_sumo.csv")).set_index("scenario").loc["pooled"]
    cd = CategoryChartData()
    cats = [("auc_none", "no context"), ("auc_env_agent", "per environment"), ("auc_rag", "retrieval"),
            ("auc_none_headway", "headway"), ("auc_rag_headway", "headway + retrieval")]
    cd.categories = [c[1] for c in cats]
    cd.add_series("AUC", [round(float(r[c[0]]), 3) for c in cats])
    x2 = X0 + 6.75
    chart_card(s, x2, 1.65, W - 6.75, 3.0, "Planted SUMO drivers, all scenarios pooled (AUC)")
    ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x2 + 0.12), Inches(2.0), Inches(W - 7.0), Inches(2.55), cd).chart
    style_chart(ch, legend=False)
    ch.category_axis.reverse_order = True
    pl = ch.plots[0]
    pl.gap_width = 40
    pl.vary_by_categories = False
    color_series(ch, [MAROON2])
    for i, pt in enumerate(pl.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(MAROON2 if i >= 3 else GREY_BAR)
    pl.has_data_labels = True
    pl.data_labels.number_format = "0.00"
    pl.data_labels.number_format_is_linked = False
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    pl.data_labels.font.size = Pt(10)
    ch.value_axis.minimum_scale = 0.6
    ch.value_axis.maximum_scale = 0.9
    ch.value_axis.major_unit = 0.1
    ch.value_axis.tick_labels.number_format = "0.00"
    ch.value_axis.tick_labels.number_format_is_linked = False
    ch.category_axis.tick_labels.font.size = Pt(9)
    f = pd.read_csv(os.path.join(DATA, "rag_ngsim_flags_headway.csv"))
    a = f[f["reference"] == "uah only"]["share_flagged_95"]
    b = f[f["reference"] == "uah + sumo"]["share_flagged_95"]
    stat(s, X0, 4.85, 5.55, f"{100 * a.min():.1f} to {100 * a.max():.0f}%", "",
         "NGSIM windows flagged with a UAH reference (headway)")
    stat(s, X0 + 5.75, 4.85, W - 5.75, f"{100 * b.min():.0f} to {100 * b.max():.0f}%", "",
         "with simulated drivers as the reference: the reference set decides", dark=True)
    text(s, X0, 5.95, W, 0.6, "On UAH (free flow, two road types) retrieval loses 0.03 AUC: there is little situation to "
         "use. Rules (speed limits, following distance by country) only explain a flag, never enter the score.",
         size=12, color=BODY)
    footer(s, n)
    return s


def slide_cutoffs(prs, crest, n):
    s = new_slide(prs)
    header(s, "Cut-offs", "Label cut-offs fitted on labelled data: 29 and 42", crest)
    c = pd.read_csv(os.path.join(DATA, "cutoffs.csv"))
    g = c.groupby(["data", "proximity", "cutoffs"])["balanced_accuracy"].mean()
    cats = [("uah", "metres", "UAH, metres"), ("uah", "headway", "UAH, headway"), ("sumo", "metres", "SUMO, metres"),
            ("sumo", "headway", "SUMO, headway")]
    cd = CategoryChartData()
    cd.categories = [k[2] for k in cats]
    cd.add_series("35 / 70 (hand set)", [round(float(g[(a, b, "35/70")]), 3) for a, b, _ in cats])
    cd.add_series("fitted", [round(float(g[(a, b, "fitted")]), 3) for a, b, _ in cats])
    chart_card(s, X0, 1.65, 6.3, 4.92, "Balanced accuracy on held-out drivers / seeds")
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(X0 + 0.15), Inches(2.05), Inches(6.0), Inches(4.4), cd).chart
    style_chart(ch)
    color_series(ch, [GREY_BAR, MAROON2])
    ch.value_axis.minimum_scale = 0.3
    ch.value_axis.maximum_scale = 0.8
    ch.value_axis.major_unit = 0.1
    ch.value_axis.tick_labels.number_format = "0.00"
    ch.value_axis.tick_labels.number_format_is_linked = False
    pl = ch.plots[0]
    pl.gap_width = 60
    pl.has_data_labels = True
    pl.data_labels.number_format = "0.00"
    pl.data_labels.number_format_is_linked = False
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    pl.data_labels.font.size = Pt(10)
    lodo = pd.read_csv(os.path.join(DATA, "uah_lodo_results.csv"))
    x2 = X0 + 6.55
    w2 = W - 6.55
    stat(s, x2, 1.65, w2, f"{lodo['original_trip_acc'].mean():.2f}", "", "UAH trip accuracy of the hand-set index on "
         "held-out drivers (0.60 with 70)")
    conf = c[(c["data"] == "sumo") & (c["proximity"] == "metres") & (c["cutoffs"] == "fitted")]["recall_aggressive"].mean()
    stat(s, x2, 2.75, w2, f"{100 * conf:.0f}%", "", "of planted aggressive drivers caught (5% with 35 / 70)")
    card(s, x2, 3.85, w2, 2.72, fill=MAROON)
    text(s, x2 + 0.25, 4.05, w2 - 0.5, 0.3, [[("Warning", {"bold": True, "color": "FFFFFF", "size": 14})]])
    u = pd.read_csv(os.path.join(DATA, "unlabelled_summary.csv")).set_index("site")
    ng = u[u.index.str.startswith("ngsim")]["share_aggressive"]
    text(s, x2 + 0.25, 4.45, w2 - 0.5, 2.0, f"With the gap in metres, {100 * ng.min():.0f} to {100 * ng.max():.0f}% of "
         "NGSIM windows now read aggressive, and most planted normal drivers in the jam. Shares of aggressive drivers "
         "in dense traffic should not be reported until proximity uses time headway.", size=12, color=PALE)
    footer(s, n)
    return s


def slide_video(prs, crest, n):
    s = new_slide(prs)
    header(s, "Replay", "Real US-101 traffic replayed in SUMO", crest)
    pst = poster(VIDEO) if os.path.exists(VIDEO) else None
    if os.path.exists(VIDEO):
        s.shapes.add_movie(VIDEO, Inches(X0), Inches(1.65), Inches(W), Inches(W * 9 / 16 * 0.62),
                           poster_frame_image=pst, mime_type="video/mp4")
    text(s, X0, 5.85, W, 0.75, "NGSIM US-101 (5 min) in SUMO, every vehicle a body ellipse plus an influence ellipse. "
         "Colours are the September scoring (context score, UAH agent); the replay will be redone with the fixed "
         "features once the proximity measure is decided.", size=12)
    footer(s, n)
    return s


def slide_next(prs, crest, n):
    s = new_slide(prs)
    header(s, "Next", "Decisions for Dr. Daher and next steps", crest)
    card(s, X0, 1.65, 5.65, 4.92, fill=MAROON)
    text(s, X0 + 0.3, 1.85, 5.1, 0.3, [[("Decisions", {"bold": True, "color": "FFFFFF", "size": 16})]])
    dec = ["Proximity: gap in metres or time headway? Every dense-traffic result favours headway; only UAH free flow "
           "prefers metres, by 0.03 AUC.",
           "Old presentation figures: redo with the fixed features, or only new ones? The September slide examples "
           "came from the wave bug.",
           "Report labels with the fitted 29 / 42 only after the proximity decision."]
    bullet_lines(s, X0 + 0.3, 2.35, 5.0, dec, color=PALE, marker=PINK, step=1.05)
    x2 = X0 + 5.9
    card(s, x2, 1.65, W - 5.9, 4.92)
    text(s, x2 + 0.3, 1.85, W - 6.5, 0.3, [[("Next steps", {"bold": True, "color": MAROON2, "size": 16})]])
    nxt = ["Full DriveDNA (465 drivers, awaiting approval) and 100-DrivingStyle (expert labels, email sent)",
           "Percentile thresholds from retrieval, with dense-traffic reference drivers",
           "SUMO layout with US-101's synchronized 48 km/h flow (speed still off)",
           "RL ego: reward for yielding at the junction",
           "Paper outline for IEEE ITSC"]
    bullet_lines(s, x2 + 0.3, 2.35, W - 6.6, nxt, step=0.7)
    footer(s, n)
    return s


def main():
    prs = Presentation(SRC)
    with zipfile.ZipFile(SRC) as z:
        tmp = tempfile.mkdtemp()
        crest = os.path.join(tmp, "crest.png")
        logo = os.path.join(tmp, "logo.png")
        open(crest, "wb").write(z.read("ppt/media/image-2-1.png"))
        open(logo, "wb").write(z.read("ppt/media/image-1-1.png"))
    for sl in prs.slides:                 # renumber the September slides
        for sh in sl.shapes:
            if sh.has_text_frame and re.fullmatch(r"\d+ / 4", sh.text_frame.text.strip()):
                r = sh.text_frame.paragraphs[0].runs[0]
                r.text = re.sub(r"/ 4$", f"/ {TOTAL}", r.text)
    divider(prs, logo)
    builders = [slide_bugs, slide_hand, slide_high, slide_planted, slide_planted_results, slide_calibration,
                slide_real, slide_labels, slide_drivedna, slide_rag, slide_cutoffs, slide_video, slide_next]
    for i, b in enumerate(builders):
        b(prs, crest, 6 + i)
    assert len(prs.slides) == TOTAL, len(prs.slides)
    prs.save(OUT)
    print("written", OUT)


if __name__ == "__main__":
    main()
